"""Sample documents for the test suite.

Everything is generated at runtime, so no binary fixtures live in git.
"""

from __future__ import annotations

import io
from pathlib import Path


def minimal_pdf(text: str) -> bytes:
    """A one-page PDF with `text` drawn on it (hand-built, no extra deps)."""
    safe = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")

    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(title: str, body: str) -> bytes:
    """Build a minimal but valid .docx with the stdlib (no python-docx)."""
    from xml.sax.saxutils import escape
    from zipfile import ZIP_DEFLATED, ZipFile

    def paragraph(text: str, style: str | None = None) -> str:
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        return f"<w:p>{ppr}<w:r><w:t>{escape(text)}</w:t></w:r></w:p>"

    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{paragraph(title, 'Heading1')}{paragraph(body)}"
        f"{paragraph('Second paragraph with some detail.')}</w:body></w:document>"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
        '<w:name w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1">'
        '<w:name w:val="Heading 1"/><w:basedOn w:val="Normal"/><w:qFormat/>'
        '<w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
        "</w:styles>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        "</Relationships>"
    )

    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("word/document.xml", document)
        archive.writestr("word/styles.xml", styles)
        archive.writestr("word/_rels/document.xml.rels", doc_rels)
    return buffer.getvalue()


def make_pptx(title: str, body: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheet: str, values: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet_obj = workbook.active
    sheet_obj.title = sheet
    for row in values:
        sheet_obj.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_samples(directory: Path) -> dict[str, Path]:
    """Create one sample per supported format; returns {label: path}."""
    directory.mkdir(parents=True, exist_ok=True)
    samples: dict[str, Path] = {}

    def put(name: str, data: bytes) -> Path:
        path = directory / name
        path.write_bytes(data)
        samples[name] = path
        return path

    put(
        "page.html",
        b"<!doctype html><html><head><title>Sample Page</title></head>"
        b"<body><h1>Welcome aboard</h1><p>This is the <strong>intro</strong>.</p>"
        b"<ul><li>alpha</li><li>beta</li></ul></body></html>",
    )
    put("notes.md", b"# Plain markdown\n\nJust *existing* markdown.\n")
    put("data.csv", b"name,score\nada,10\ngrace,20\n")
    put("readme.txt", b"Plain text line one.\nPlain text line two.\n")
    put("report.pdf", minimal_pdf("PDF body text for conversion testing."))
    put("deck.pptx", make_pptx("Slide Title", "Bullet one on the slide"))
    put(
        "book.docx",
        make_docx("Contract Heading", "The obligor shall deliver the goods."),
    )
    put(
        "figures.xlsx",
        make_xlsx("Numbers", [["item", "qty"], ["widgets", 42]]),
    )
    put("junk.dat", bytes(range(256)) * 8)  # not a document
    return samples


__all__ = [
    "make_docx",
    "make_pptx",
    "make_xlsx",
    "minimal_pdf",
    "write_samples",
]
